import http.server
import socketserver
import json
import os
import openpyxl
import re
import webbrowser
import sys

PORT = 8000
SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(SCRIPTS_DIR)
WEB_DIR = os.path.join(BASE_DIR, 'web')
DOCS_DIR = os.path.join(BASE_DIR, 'docs')
XLSX_PATH = os.path.join(BASE_DIR, 'LeetCode_DSA_Tracker.xlsx')
PROBLEMS_JSON = os.path.join(WEB_DIR, 'problems.json')
PROBLEMS_JS = os.path.join(WEB_DIR, 'problems_data.js')

class DSARevisionHandler(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=WEB_DIR, **kwargs)

    def do_GET(self):
        if self.path == '/api/load':
            self.send_response(200)
            self.send_header('Content-Type', 'application/json')
            self.send_header('Access-Control-Allow-Origin', '*')
            self.end_headers()
            if os.path.exists(PROBLEMS_JSON):
                with open(PROBLEMS_JSON, 'r', encoding='utf-8') as f:
                    self.wfile.write(f.read().encode('utf-8'))
            else:
                self.wfile.write(b'[]')
            return
        elif self.path == '/api/health':
            self.send_response(200)
            self.send_header('Content-Type', 'application/json')
            self.send_header('Access-Control-Allow-Origin', '*')
            self.end_headers()
            self.wfile.write(b'{"status": "ok", "autoSave": true}')
            return
        return super().do_GET()

    def do_POST(self):
        if self.path == '/api/save':
            content_length = int(self.headers.get('Content-Length', 0))
            post_data = self.rfile.read(content_length)
            try:
                data = json.loads(post_data.decode('utf-8'))
                problems = data.get('problems', data) if isinstance(data, dict) else data

                # 1. Save to problems.json (web and docs)
                for folder in [WEB_DIR, DOCS_DIR]:
                    if os.path.exists(folder):
                        with open(os.path.join(folder, 'problems.json'), 'w', encoding='utf-8') as f:
                            json.dump(problems, f, indent=2, ensure_ascii=False)

                # 2. Save to problems_data.js (web and docs)
                js_content = f"// Auto-saved from DSA Revision App\nwindow.INITIAL_PROBLEMS = {json.dumps(problems, indent=2, ensure_ascii=False)};\n"
                for folder in [WEB_DIR, DOCS_DIR]:
                    if os.path.exists(folder):
                        with open(os.path.join(folder, 'problems_data.js'), 'w', encoding='utf-8') as f:
                            f.write(js_content)

                # 3. Save directly into LeetCode_DSA_Tracker.xlsx
                excel_updated = self.update_excel(problems)

                self.send_response(200)
                self.send_header('Content-Type', 'application/json')
                self.send_header('Access-Control-Allow-Origin', '*')
                self.end_headers()
                resp = {
                    "success": True,
                    "savedJson": True,
                    "savedExcel": excel_updated,
                    "count": len(problems)
                }
                self.wfile.write(json.dumps(resp).encode('utf-8'))
                print(f"[Auto-Save] Successfully saved {len(problems)} problems to disk and Excel!")
            except Exception as e:
                print(f"[Auto-Save Error]: {e}")
                self.send_response(500)
                self.send_header('Content-Type', 'application/json')
                self.send_header('Access-Control-Allow-Origin', '*')
                self.end_headers()
                self.wfile.write(json.dumps({"success": False, "error": str(e)}).encode('utf-8'))
            return

        self.send_response(404)
        self.end_headers()

    def do_OPTIONS(self):
        self.send_response(200)
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.end_headers()

    def update_excel(self, problems):
        if not os.path.exists(XLSX_PATH):
            return False

        problem_map = {}
        for p in problems:
            if 'name' in p:
                problem_map[p['name'].strip().lower()] = p

        try:
            wb = openpyxl.load_workbook(XLSX_PATH)
            ws = wb.active

            for row_idx in range(2, ws.max_row + 1):
                cell_name = ws.cell(row=row_idx, column=4).value
                if not cell_name:
                    continue

                name_key = str(cell_name).strip().lower()
                if name_key in problem_map:
                    p_data = problem_map[name_key]
                    if p_data.get('solvedAlone'):
                        ws.cell(row=row_idx, column=8, value='YES')
                    if p_data.get('dateSolved'):
                        ws.cell(row=row_idx, column=9, value=p_data['dateSolved'])
                    if p_data.get('rev1'):
                        ws.cell(row=row_idx, column=10, value=p_data['rev1'])
                    if p_data.get('rev2'):
                        ws.cell(row=row_idx, column=11, value=p_data['rev2'])
                    if p_data.get('rev3'):
                        ws.cell(row=row_idx, column=12, value=p_data['rev3'])

            wb.save(XLSX_PATH)
            return True
        except PermissionError:
            print("[Excel Warning] Could not write to LeetCode_DSA_Tracker.xlsx because it is currently open in Excel. Changes were saved to JSON/JS.")
            return False
        except Exception as e:
            print(f"[Excel Error]: {e}")
            return False

def start_server():
    socketserver.TCPServer.allow_reuse_address = True
    try:
        with socketserver.TCPServer(("", PORT), DSARevisionHandler) as httpd:
            print(f"============================================================")
            print(f"  DSA Revision Server running at: http://localhost:{PORT}")
            print(f"  Auto-Save Active: Every revision is saved directly to disk!")
            print(f"============================================================")
            
            # Auto open browser
            if '--no-browser' not in sys.argv:
                webbrowser.open(f"http://localhost:{PORT}")
                
            httpd.serve_forever()
    except OSError as e:
        if "address already in use" in str(e).lower() or e.errno == 98 or e.errno == 10048:
            print(f"[Notice] Port {PORT} is already in use.")
            if '--no-browser' not in sys.argv:
                webbrowser.open(f"http://localhost:{PORT}")
        else:
            raise

if __name__ == '__main__':
    start_server()
